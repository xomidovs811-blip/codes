"""
Builds the "Zayavkalar hisoboti" workbook for whatever filters were applied
on the table page (the same list the user sees), sent to the Telegram
channel on demand. It has two sheets, both built from each zayavka's latest
version:

  - "Hisobot": one row per zayavka (header fields, totals, sign-off names);
  - "Qatorlar": every item row of every zayavka, so the full contents of the
    forms are in the file.

Every report states, at the top of the sheet, when it was generated, who
sent it and which filters were used, and the file name / Telegram caption
carry the same details - so two exports made with different criteria (or at
different times) can always be told apart.
"""
import html
import re
from datetime import date, datetime
from pathlib import Path
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "1F4E78"
NAVY_TEXT = "17365D"
LABEL_BG = "EEF5FB"
TOTAL_GREEN_BG = "D9EAD3"
TOTAL_GREEN_TEXT = "274E13"

THIN = Side(style="thin", color="C9D6E3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")

# (header, width)
COLUMNS = [
    ("Т/р", 6),
    ("Номер заявки", 18),
    ("Дата", 12),
    ("Название объекта", 26),
    ("От кого", 24),
    ("Кому", 18),
    ("Вид оплаты", 15),
    ("Qator", 8),
    ("Общая сумма", 17),
    ("Аванс получил", 17),
    ("Остатка", 17),
    ("Oxirgi yuborilgan", 18),
    ("Versiya", 9),
    ("Инспектор", 22),
    ("Кассир", 22),
    ("Техданзор", 22),
]
HEADER_ROW = 9

# Second sheet "Qatorlar": every item row of every zayavka's latest version,
# flat (one row per item), so nothing on the forms is left out of the report.
ITEM_COLUMNS = [
    ("Т/р", 6),
    ("Номер заявки", 18),
    ("Дата", 12),
    ("Название объекта", 24),
    ("От кого", 22),
    ("Вид оплаты", 14),
    ("№", 5),
    ("Наименование товара", 42),
    ("Поставщик", 22),
    ("Блок", 8),
    ("Этаж", 8),
    ("Ед.изм", 9),
    ("Кол-во", 10),
    ("Цена", 14),
    ("Общая сумма", 16),
    ("Аванс получил", 16),
    ("Остатка", 16),
    ("Выполняемая работа (сметная группа)", 40),
    ("Комментарие", 26),
]


def describe_filters(filters: dict) -> dict:
    """Human-readable filter values, shared by the sheet, file name and caption."""
    d_from: Optional[date] = filters.get("date_from")
    d_to: Optional[date] = filters.get("date_to")
    if d_from and d_to:
        period = f"{d_from:%d.%m.%Y} - {d_to:%d.%m.%Y}"
    elif d_from:
        period = f"{d_from:%d.%m.%Y} dan"
    elif d_to:
        period = f"{d_to:%d.%m.%Y} gacha"
    else:
        period = "Barchasi"
    return {
        "object": filters.get("object_name") or "Barchasi",
        "from": filters.get("from_whom") or "Barchasi",
        "period": period,
    }


def report_filename(filters: dict, generated_at: datetime) -> str:
    f = describe_filters(filters)
    parts = [f"Hisobot_{generated_at:%Y-%m-%d_%H-%M}"]
    if filters.get("object_name"):
        parts.append(f["object"])
    if filters.get("from_whom"):
        parts.append(f["from"])
    if filters.get("date_from") or filters.get("date_to"):
        parts.append(f["period"].replace(" - ", "_").replace(" ", "_"))
    name = "_".join(parts)
    name = re.sub(r'[\\/:*?"<>|\s]+', "_", name)[:120].strip("_")
    return name + ".xlsx"


def _row_sums(z: dict) -> tuple:
    items = z.get("items") or []
    return (
        sum((i.get("total") or 0) for i in items),
        sum((i.get("advance") or 0) for i in items),
        sum((i.get("remainder") or 0) for i in items),
    )


def _item_count(rows: List[dict]) -> int:
    return sum(len(z.get("items") or []) for z in rows)


def totals(rows: List[dict]) -> tuple:
    t = a = r = 0
    for z in rows:
        zt, za, zr = _row_sums(z)
        t, a, r = t + zt, a + za, r + zr
    return t, a, r


def build_report_caption(
    rows: List[dict], filters: dict, generated_at: datetime, requested_by: str
) -> str:
    """Telegram caption (HTML): the same identifying details as the sheet's header block."""
    f = describe_filters(filters)
    t, a, r = totals(rows)
    money = lambda n: f"{n:,.0f}".replace(",", " ")
    e = html.escape
    return (
        "📊 <b>Zayavkalar hisoboti</b>\n\n"
        f"🕒 <b>Yaratildi:</b> {generated_at:%d.%m.%Y %H:%M}\n"
        f"👤 <b>Yubordi:</b> {e(requested_by or '—')}\n"
        f"🏢 <b>Obyekt:</b> {e(f['object'])}\n"
        f"📤 <b>От кого:</b> {e(f['from'])}\n"
        f"📅 <b>Sana:</b> {e(f['period'])}\n"
        f"📄 <b>Zayavkalar:</b> {len(rows)} ta\n"
        f"🧾 <b>Qatorlar (tovar/xizmat):</b> {_item_count(rows)} ta\n\n"
        f"💰 <b>Общая сумма:</b> {money(t)}\n"
        f"💵 <b>Аванс:</b> {money(a)}\n"
        f"✅ <b>Остатка:</b> {money(r)}"
    )


def build_report_workbook(
    rows: List[dict], filters: dict, generated_at: datetime, requested_by: str
) -> Workbook:
    """
    rows: zayavka dicts (Zayavka.to_dict()) plus "version_count" and
    "last_saved_at" (a naive datetime already in local/Tashkent time, or None).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Hisobot"
    ws.sheet_view.showGridLines = False

    for idx, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    last_col = len(COLUMNS)

    # --- title banner ---
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title = ws.cell(row=1, column=1, value="ZAYAVKALAR HISOBOTI")
    title.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.alignment = CENTER
    ws.row_dimensions[1].height = 28

    # --- what makes this report distinct: when / who / which filters ---
    f = describe_filters(filters)
    info = [
        ("Hisobot yaratilgan:", generated_at.strftime("%d.%m.%Y %H:%M")),
        ("Kim yubordi:", requested_by or "—"),
        ("Название объекта:", f["object"]),
        ("От кого:", f["from"]),
        ("Sana oralig'i (Дата):", f["period"]),
        ("Zayavkalar soni:", f"{len(rows)} ta zayavka, {_item_count(rows)} ta qator (\"Qatorlar\" varag'ida)"),
    ]
    for offset, (label, value) in enumerate(info):
        r = 2 + offset
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        lab = ws.cell(row=r, column=1, value=label)
        lab.font = Font(name="Calibri", size=11, bold=True, color=NAVY_TEXT)
        lab.fill = PatternFill("solid", fgColor=LABEL_BG)
        lab.alignment = LEFT
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
        val = ws.cell(row=r, column=4, value=value)
        val.font = Font(name="Calibri", size=11, bold=True, color=NAVY_TEXT)
        val.alignment = LEFT
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = BORDER
        ws.row_dimensions[r].height = 20

    # --- table header ---
    ws.row_dimensions[HEADER_ROW].height = 32
    for idx, (label, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=idx, value=label)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = CENTER
        cell.border = BORDER

    # --- rows ---
    first = HEADER_ROW + 1
    r = first
    for n, z in enumerate(rows, start=1):
        total, advance, remainder = _row_sums(z)
        values = [
            n,
            z.get("number"),
            z.get("date"),
            z.get("object_name"),
            z.get("from_whom"),
            z.get("to_whom"),
            z.get("payment_type"),
            len(z.get("items") or []),
            total,
            advance,
            remainder,
            z.get("last_saved_at"),
            z.get("version_count") or 1,
            z.get("inspector") or "",
            z.get("cashier") or "",
            z.get("tech_supervisor") or "",
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            cell.font = Font(name="Calibri", size=11, bold=(c == 2))
            cell.alignment = LEFT
            if c in (1, 8, 13):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if c in (9, 10, 11):
                cell.alignment = RIGHT
                cell.number_format = "#,##0"
            if c == 3:
                cell.number_format = "DD.MM.YYYY"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if c == 12:
                cell.number_format = "DD.MM.YYYY HH:MM"
                cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=11).fill = PatternFill("solid", fgColor=LABEL_BG)
        ws.row_dimensions[r].height = 21
        r += 1
    last = r - 1

    # --- totals row (live formulas over the rows above) ---
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    label = ws.cell(row=r, column=1, value="JAMI")
    label.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    label.alignment = LEFT
    for c in range(1, last_col + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = BORDER
        if c != 11:
            cell.fill = PatternFill("solid", fgColor=NAVY)
    for c in (9, 10, 11):
        col = get_column_letter(c)
        cell = ws.cell(row=r, column=c, value=f"=SUM({col}{first}:{col}{last})" if rows else 0)
        cell.number_format = "#,##0"
        cell.alignment = RIGHT
        cell.font = Font(
            name="Calibri", size=12, bold=True,
            color=TOTAL_GREEN_TEXT if c == 11 else "FFFFFF",
        )
        if c == 11:
            cell.fill = PatternFill("solid", fgColor=TOTAL_GREEN_BG)
    ws.row_dimensions[r].height = 24

    ws.freeze_panes = ws.cell(row=first, column=1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{get_column_letter(last_col)}{r}"

    _build_items_sheet(wb.create_sheet("Qatorlar"), rows)
    return wb


def _build_items_sheet(ws, rows: List[dict]) -> None:
    """
    One row per item of each zayavka's latest version. Общая сумма / Остатка
    are live formulas (like in each zayavka's own Excel); the total row uses
    SUBTOTAL so it follows whatever the reader filters. Zayavkas are banded
    (alternate shading) so their rows are easy to tell apart.
    """
    ws.sheet_view.showGridLines = False
    for idx, (_, width) in enumerate(ITEM_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ncols = len(ITEM_COLUMNS)

    ws.row_dimensions[1].height = 34
    for idx, (label, _) in enumerate(ITEM_COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=label)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = CENTER
        cell.border = BORDER

    band = PatternFill("solid", fgColor=LABEL_BG)
    r = 2
    n = 0
    for zi, z in enumerate(rows):
        items = z.get("items") or []
        for item in items:
            n += 1
            values = [
                n,
                z.get("number"),
                z.get("date"),
                z.get("object_name"),
                z.get("from_whom"),
                z.get("payment_type"),
                item.get("row_no"),
                item.get("product_name") or "",
                item.get("supplier") or "",
                item.get("block") or "",
                item.get("floor") or "",
                item.get("unit") or "",
                item.get("qty"),
                item.get("price"),
                f"=M{r}*N{r}",
                item.get("advance"),
                f"=O{r}-P{r}",
                item.get("work_type") or "",
                item.get("comment") or "",
            ]
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.border = BORDER
                cell.font = Font(name="Calibri", size=11, bold=(c == 2))
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=c in (8, 18, 19))
                if c in (1, 7, 10, 11, 12):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                if c == 3:
                    cell.number_format = "DD.MM.YYYY"
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                if c in (13, 14, 15, 16, 17):
                    cell.alignment = RIGHT
                    cell.number_format = "#,##0.##" if c == 13 else "#,##0"
                if zi % 2 == 0:
                    cell.fill = band
            r += 1
    last = r - 1

    if last >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{last}"
    ws.freeze_panes = "C2"

    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14)
    label = ws.cell(row=r, column=1, value="JAMI (filtrga qarab)")
    label.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    label.alignment = LEFT
    for c in range(1, ncols + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = BORDER
        if c != 17:
            cell.fill = PatternFill("solid", fgColor=NAVY)
    for c in (15, 16, 17):
        col = get_column_letter(c)
        cell = ws.cell(row=r, column=c, value=f"=SUBTOTAL(109,{col}2:{col}{last})" if last >= 2 else 0)
        cell.number_format = "#,##0"
        cell.alignment = RIGHT
        cell.font = Font(
            name="Calibri", size=12, bold=True,
            color=TOTAL_GREEN_TEXT if c == 17 else "FFFFFF",
        )
        if c == 17:
            cell.fill = PatternFill("solid", fgColor=TOTAL_GREEN_BG)
    ws.row_dimensions[r].height = 24

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def save_report(
    rows: List[dict], filters: dict, generated_at: datetime, requested_by: str, out_dir: str
) -> str:
    wb = build_report_workbook(rows, filters, generated_at, requested_by)
    path = str(Path(out_dir) / report_filename(filters, generated_at))
    wb.save(path)
    return path
